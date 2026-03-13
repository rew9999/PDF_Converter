"""PDF変換ツール - メインアプリケーション"""

import json
import os
import sys
import shutil
import tempfile
import uuid
import webbrowser
from pathlib import Path
from threading import Timer

import io

import fitz  # PyMuPDF
import pytesseract

# Windows環境でTesseractのパスを自動検出
if sys.platform == "win32" and not shutil.which("tesseract"):
    _tesseract_candidates = [
        r"C:\Program Files\Tesseract-OCR\tesseract.exe",
        r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
    ]
    for _path in _tesseract_candidates:
        if os.path.isfile(_path):
            pytesseract.pytesseract.tesseract_cmd = _path
            break
from docx import Document
from docx.shared import Inches, Pt
from fastapi import FastAPI, File, HTTPException, Request, UploadFile
from fastapi.responses import FileResponse, HTMLResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XlImage
from PIL import Image
from pydantic import BaseModel

# 対応フォーマット
IMAGE_FORMATS = ("png", "jpg")
DOCUMENT_FORMATS = ("docx", "xlsx")
ALL_FORMATS = IMAGE_FORMATS + DOCUMENT_FORMATS

# パス設定（PyInstaller対応）
if getattr(sys, "frozen", False):
    BASE_DIR = Path(sys._MEIPASS)
    APP_DIR = Path(sys.executable).parent
else:
    BASE_DIR = Path(__file__).resolve().parent
    APP_DIR = BASE_DIR

CONFIG_PATH = APP_DIR / "config.json"
TEMP_DIR = APP_DIR / "temp"
TEMP_DIR.mkdir(exist_ok=True)

# 設定読み込み
DEFAULT_CONFIG = {
    "default_output_dir": "",
    "default_format": "png",
    "default_dpi": 200,
    "max_file_size_mb": 100,
    "temp_dir": "./temp",
    "port": 5000,
    "auto_open_browser": True,
}


def load_config() -> dict:
    if CONFIG_PATH.exists():
        with open(CONFIG_PATH, "r", encoding="utf-8") as f:
            user_config = json.load(f)
        config = {**DEFAULT_CONFIG, **user_config}
    else:
        config = DEFAULT_CONFIG.copy()
        with open(CONFIG_PATH, "w", encoding="utf-8") as f:
            json.dump(config, f, indent=2, ensure_ascii=False)
    return config


config = load_config()

# FastAPIアプリ初期化
app = FastAPI(title="PDF変換ツール", version="1.0.0")
app.mount("/static", StaticFiles(directory=str(BASE_DIR / "static")), name="static")
templates = Jinja2Templates(directory=str(BASE_DIR / "templates"))

# アップロード済みファイルの管理 {file_id: {"path": Path, "filename": str}}
uploaded_files: dict[str, dict] = {}
# 変換ジョブの管理 {job_id: {"status": str, "progress": int, ...}}
conversion_jobs: dict[str, dict] = {}


# ---------- リクエストモデル ----------


class ConvertRequest(BaseModel):
    file_ids: list[str]
    format: str = "png"
    dpi: int = 200
    output_dir: str = ""


# ---------- ユーティリティ ----------


def cleanup_temp_file(file_path: Path):
    """一時ファイルを安全に削除"""
    try:
        if file_path.exists():
            file_path.unlink()
    except OSError:
        pass


def sanitize_filename(filename: str) -> str:
    """ファイル名のサニタイズ"""
    # パストラバーサル対策
    filename = Path(filename).name
    # 危険な文字を除去
    keepchars = (" ", ".", "_", "-")
    return "".join(c for c in filename if c.isalnum() or c in keepchars).strip()


def convert_pdf_to_images(
    pdf_path: Path,
    output_dir: Path,
    fmt: str,
    dpi: int,
    job_id: str,
    file_index: int,
    total_files: int,
    original_filename: str = "",
) -> list[Path]:
    """PDFを画像に変換する"""
    output_files = []
    doc = fitz.open(str(pdf_path))
    total_pages = len(doc)
    # 元のPDFファイル名から拡張子を除いた名前を使用
    stem = Path(original_filename).stem if original_filename else pdf_path.stem

    for page_num in range(total_pages):
        page = doc[page_num]
        # DPIに基づくズーム倍率（デフォルト72dpi基準）
        zoom = dpi / 72
        matrix = fitz.Matrix(zoom, zoom)
        pix = page.get_pixmap(matrix=matrix)

        # Pillow Imageに変換
        img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)

        # ファイル名: 元のPDF名_連番.拡張子（1始まり）
        out_name = f"{stem}_{page_num + 1}.{fmt}"
        out_path = output_dir / out_name

        # 同名ファイルが存在する場合はサフィックス追加
        counter = 1
        while out_path.exists():
            out_name = f"{stem}_{page_num + 1}({counter}).{fmt}"
            out_path = output_dir / out_name
            counter += 1

        if fmt == "jpg":
            img.save(str(out_path), "JPEG", quality=95)
        else:
            img.save(str(out_path), "PNG")

        output_files.append(out_path)

        # 進捗更新
        if job_id in conversion_jobs:
            file_progress = ((page_num + 1) / total_pages) * 100
            overall_progress = (
                (file_index * 100 + file_progress) / total_files
            )
            conversion_jobs[job_id]["progress"] = int(overall_progress)
            conversion_jobs[job_id]["current_page"] = page_num + 1
            conversion_jobs[job_id]["total_pages"] = total_pages

    doc.close()
    return output_files


def _extract_page_text_blocks(page) -> list[str]:
    """ページからテキストを抽出する（複数の方法を試行）"""
    # 方法1: blocksモードで抽出
    blocks = page.get_text("blocks")
    text_blocks = []
    for b in sorted(blocks, key=lambda b: (b[1], b[0])):
        if b[6] == 0:  # テキストブロック
            text = b[4].strip()
            if text:
                text_blocks.append(text)

    if text_blocks:
        return text_blocks

    # 方法2: プレーンテキストで抽出
    plain = page.get_text("text").strip()
    if plain:
        return [line for line in plain.split("\n") if line.strip()]

    return []


def _render_page_to_image_bytes(page, dpi: int = 200) -> bytes:
    """ページを画像としてレンダリングし、PNGバイト列を返す"""
    zoom = dpi / 72
    matrix = fitz.Matrix(zoom, zoom)
    pix = page.get_pixmap(matrix=matrix)
    return pix.tobytes("png")


def _get_ocr_lang() -> str:
    """利用可能なTesseract言語を検出してOCR用言語文字列を返す"""
    try:
        langs = pytesseract.get_languages(config="")
        available = set(langs)
        # 日本語+英語を優先、なければ英語のみ、それもなければデフォルト
        if "jpn" in available and "eng" in available:
            return "jpn+eng"
        if "jpn" in available:
            return "jpn"
        if "eng" in available:
            return "eng"
        return langs[0] if langs else "eng"
    except Exception:
        return "eng"


def _ocr_page(page, dpi: int = 300) -> str:
    """ページを画像化してOCRでテキストを抽出する"""
    try:
        img_bytes = _render_page_to_image_bytes(page, dpi=dpi)
        pil_image = Image.open(io.BytesIO(img_bytes))
        lang = _get_ocr_lang()
        text = pytesseract.image_to_string(pil_image, lang=lang)
        return text.strip()
    except Exception:
        return ""


def convert_pdf_to_docx(
    pdf_path: Path,
    output_dir: Path,
    job_id: str,
    file_index: int,
    total_files: int,
    original_filename: str = "",
) -> list[Path]:
    """PDFをWord(.docx)に変換する"""
    stem = Path(original_filename).stem if original_filename else pdf_path.stem
    out_name = f"{stem}.docx"
    out_path = output_dir / out_name

    # 同名ファイルが存在する場合はサフィックス追加
    counter = 1
    while out_path.exists():
        out_name = f"{stem}({counter}).docx"
        out_path = output_dir / out_name
        counter += 1

    pdf_doc = fitz.open(str(pdf_path))
    total_pages = len(pdf_doc)
    word_doc = Document()
    temp_images = []

    for page_num in range(total_pages):
        page = pdf_doc[page_num]

        if page_num > 0:
            word_doc.add_page_break()

        # テキスト抽出を試行
        text_blocks = _extract_page_text_blocks(page)

        if text_blocks:
            # テキストが取得できた場合はテキストとして挿入
            for text in text_blocks:
                word_doc.add_paragraph(text)
        else:
            # テキストが取得できない場合（スキャンPDF等）→ OCRで文字認識
            ocr_text = _ocr_page(page)
            if ocr_text:
                for line in ocr_text.split("\n"):
                    if line.strip():
                        word_doc.add_paragraph(line.strip())
            else:
                # OCRでも取得できない場合は画像として挿入
                img_bytes = _render_page_to_image_bytes(page)
                img_stream = io.BytesIO(img_bytes)
                word_doc.add_paragraph()
                word_doc.add_picture(img_stream, width=Inches(6))

        # 進捗更新
        if job_id in conversion_jobs:
            file_progress = ((page_num + 1) / total_pages) * 100
            overall_progress = (file_index * 100 + file_progress) / total_files
            conversion_jobs[job_id]["progress"] = int(overall_progress)
            conversion_jobs[job_id]["current_page"] = page_num + 1
            conversion_jobs[job_id]["total_pages"] = total_pages

    word_doc.save(str(out_path))
    pdf_doc.close()
    return [out_path]


def convert_pdf_to_xlsx(
    pdf_path: Path,
    output_dir: Path,
    job_id: str,
    file_index: int,
    total_files: int,
    original_filename: str = "",
) -> list[Path]:
    """PDFをExcel(.xlsx)に変換する"""
    stem = Path(original_filename).stem if original_filename else pdf_path.stem
    out_name = f"{stem}.xlsx"
    out_path = output_dir / out_name

    # 同名ファイルが存在する場合はサフィックス追加
    counter = 1
    while out_path.exists():
        out_name = f"{stem}({counter}).xlsx"
        out_path = output_dir / out_name
        counter += 1

    pdf_doc = fitz.open(str(pdf_path))
    total_pages = len(pdf_doc)
    wb = Workbook()
    wb.remove(wb.active)  # デフォルトシートを削除

    for page_num in range(total_pages):
        page = pdf_doc[page_num]
        ws = wb.create_sheet(title=f"Page{page_num + 1}")

        # wordsモードで個々の単語を座標付きで抽出
        # words: (x0, y0, x1, y1, "word", block_no, line_no, word_no)
        words = page.get_text("words")
        words = [w for w in words if w[4].strip()]

        if words:
            # Y座標でグループ化して行に、X座標でソートして列に配置
            words.sort(key=lambda w: (w[1], w[0]))
            rows = []
            current_row = [words[0]]
            for w in words[1:]:
                # Y座標の差が小さければ同じ行（フォントサイズの半分程度を閾値に）
                if abs(w[1] - current_row[0][1]) < 10:
                    current_row.append(w)
                else:
                    rows.append(current_row)
                    current_row = [w]
            rows.append(current_row)

            for row_idx, row_words in enumerate(rows, start=1):
                row_words.sort(key=lambda w: w[0])

                # X座標の間隔が大きい箇所で列を分割
                columns = [[row_words[0]]]
                for w in row_words[1:]:
                    prev_end = columns[-1][-1][2]  # 前の単語の右端x1
                    gap = w[0] - prev_end  # 現在の単語の左端x0との差
                    if gap > 20:  # 間隔が大きければ新しい列
                        columns.append([w])
                    else:
                        columns[-1].append(w)

                for col_idx, col_words in enumerate(columns, start=1):
                    text = " ".join(w[4] for w in col_words).strip()
                    # 数値として解釈できる場合は数値で入力
                    try:
                        value = float(text.replace(",", ""))
                        if value == int(value):
                            value = int(value)
                        ws.cell(row=row_idx, column=col_idx, value=value)
                    except ValueError:
                        ws.cell(row=row_idx, column=col_idx, value=text)
        else:
            # テキストが取得できない場合: OCRで文字認識
            ocr_text = _ocr_page(page)
            if ocr_text:
                lines = [l for l in ocr_text.split("\n") if l.strip()]
                for row_idx, line in enumerate(lines, start=1):
                    # タブやスペース区切りで列を分割
                    parts = line.split("\t") if "\t" in line else line.split()
                    for col_idx, part in enumerate(parts, start=1):
                        text = part.strip()
                        if not text:
                            continue
                        try:
                            value = float(text.replace(",", ""))
                            if value == int(value):
                                value = int(value)
                            ws.cell(row=row_idx, column=col_idx, value=value)
                        except ValueError:
                            ws.cell(row=row_idx, column=col_idx, value=text)
            else:
                # OCRでもテキストが取得できない場合は画像を埋め込み
                img_bytes = _render_page_to_image_bytes(page)
                img_stream = io.BytesIO(img_bytes)
                xl_img = XlImage(img_stream)
                xl_img.width = 600
                xl_img.height = int(600 * page.rect.height / page.rect.width)
                ws.add_image(xl_img, "A1")

        # 進捗更新
        if job_id in conversion_jobs:
            file_progress = ((page_num + 1) / total_pages) * 100
            overall_progress = (file_index * 100 + file_progress) / total_files
            conversion_jobs[job_id]["progress"] = int(overall_progress)
            conversion_jobs[job_id]["current_page"] = page_num + 1
            conversion_jobs[job_id]["total_pages"] = total_pages

    wb.save(str(out_path))
    pdf_doc.close()
    return [out_path]


# ---------- エンドポイント ----------


@app.get("/", response_class=HTMLResponse)
async def index(request: Request):
    """メイン画面"""
    return templates.TemplateResponse("index.html", {
        "request": request,
        "config": config,
    })


@app.post("/upload")
async def upload_file(file: UploadFile = File(...)):
    """PDFファイルのアップロード"""
    # ファイル形式チェック
    if not file.filename or not file.filename.lower().endswith(".pdf"):
        raise HTTPException(status_code=400, detail="PDFファイルのみアップロード可能です")

    # ファイルサイズチェック
    content = await file.read()
    max_size = config["max_file_size_mb"] * 1024 * 1024
    if len(content) > max_size:
        raise HTTPException(
            status_code=400,
            detail=f"ファイルサイズが上限（{config['max_file_size_mb']}MB）を超えています",
        )

    # PDFとして有効か検証
    try:
        doc = fitz.open(stream=content, filetype="pdf")
        page_count = len(doc)
        doc.close()
    except Exception:
        raise HTTPException(status_code=400, detail="無効なPDFファイルです")

    # 一時ファイルとして保存
    file_id = str(uuid.uuid4())
    safe_name = sanitize_filename(file.filename)
    temp_path = TEMP_DIR / f"{file_id}_{safe_name}"

    with open(temp_path, "wb") as f:
        f.write(content)

    uploaded_files[file_id] = {
        "path": temp_path,
        "filename": safe_name,
        "size": len(content),
        "page_count": page_count,
    }

    return {
        "file_id": file_id,
        "filename": safe_name,
        "size": len(content),
        "page_count": page_count,
    }


@app.post("/convert")
async def convert(req: ConvertRequest):
    """変換実行"""
    # バリデーション
    if req.format not in ALL_FORMATS:
        raise HTTPException(status_code=400, detail=f"対応していない形式です（{', '.join(ALL_FORMATS)}）")
    if req.format in IMAGE_FORMATS and req.dpi not in (150, 200, 300):
        raise HTTPException(status_code=400, detail="対応していない解像度です（150, 200, 300）")
    if not req.file_ids:
        raise HTTPException(status_code=400, detail="変換するファイルが選択されていません")

    # アップロード済みファイルの確認
    for fid in req.file_ids:
        if fid not in uploaded_files:
            raise HTTPException(status_code=404, detail=f"ファイルが見つかりません: {fid}")

    # 出力先ディレクトリの決定
    if req.output_dir:
        output_dir = Path(req.output_dir)
    elif config["default_output_dir"]:
        output_dir = Path(config["default_output_dir"])
    else:
        # デフォルト: ダウンロードフォルダまたはデスクトップ
        home = Path.home()
        downloads = home / "Downloads"
        if downloads.exists():
            output_dir = downloads / "PDF変換出力"
        else:
            output_dir = home / "Desktop" / "PDF変換出力"

    output_dir.mkdir(parents=True, exist_ok=True)

    # ジョブ作成
    job_id = str(uuid.uuid4())
    conversion_jobs[job_id] = {
        "status": "processing",
        "progress": 0,
        "current_file": "",
        "current_page": 0,
        "total_pages": 0,
        "completed": 0,
        "total": len(req.file_ids),
        "output_dir": str(output_dir),
        "output_files": [],
        "errors": [],
    }

    # 変換処理
    total_files = len(req.file_ids)
    for i, file_id in enumerate(req.file_ids):
        file_info = uploaded_files[file_id]
        conversion_jobs[job_id]["current_file"] = file_info["filename"]

        try:
            if req.format in IMAGE_FORMATS:
                result_files = convert_pdf_to_images(
                    pdf_path=file_info["path"],
                    output_dir=output_dir,
                    fmt=req.format,
                    dpi=req.dpi,
                    job_id=job_id,
                    file_index=i,
                    total_files=total_files,
                    original_filename=file_info["filename"],
                )
            elif req.format == "docx":
                result_files = convert_pdf_to_docx(
                    pdf_path=file_info["path"],
                    output_dir=output_dir,
                    job_id=job_id,
                    file_index=i,
                    total_files=total_files,
                    original_filename=file_info["filename"],
                )
            elif req.format == "xlsx":
                result_files = convert_pdf_to_xlsx(
                    pdf_path=file_info["path"],
                    output_dir=output_dir,
                    job_id=job_id,
                    file_index=i,
                    total_files=total_files,
                    original_filename=file_info["filename"],
                )
            conversion_jobs[job_id]["output_files"].extend(
                [str(f) for f in result_files]
            )
            conversion_jobs[job_id]["completed"] = i + 1
        except Exception as e:
            conversion_jobs[job_id]["errors"].append(
                {"file": file_info["filename"], "error": str(e)}
            )

        # アップロード一時ファイルを削除
        cleanup_temp_file(file_info["path"])
        del uploaded_files[file_id]

    # ジョブ完了
    if conversion_jobs[job_id]["errors"]:
        conversion_jobs[job_id]["status"] = "completed_with_errors"
    else:
        conversion_jobs[job_id]["status"] = "completed"
    conversion_jobs[job_id]["progress"] = 100

    return {"job_id": job_id}


@app.get("/status/{job_id}")
async def get_status(job_id: str):
    """変換進捗の取得"""
    if job_id not in conversion_jobs:
        raise HTTPException(status_code=404, detail="ジョブが見つかりません")
    return conversion_jobs[job_id]


@app.delete("/upload/{file_id}")
async def delete_uploaded_file(file_id: str):
    """アップロード済みファイルの削除"""
    if file_id not in uploaded_files:
        raise HTTPException(status_code=404, detail="ファイルが見つかりません")

    cleanup_temp_file(uploaded_files[file_id]["path"])
    del uploaded_files[file_id]
    return {"message": "削除しました"}


@app.get("/config")
async def get_config():
    """現在の設定を取得"""
    return config


@app.get("/ocr-status")
async def ocr_status():
    """OCRの利用可能状況を返す"""
    result = {"tesseract_available": False, "languages": [], "jpn_available": False}
    try:
        langs = pytesseract.get_languages(config="")
        result["tesseract_available"] = True
        result["languages"] = langs
        result["jpn_available"] = "jpn" in langs
    except Exception:
        pass
    return result


@app.post("/ocr-install-jpn")
async def ocr_install_jpn():
    """日本語OCRデータをダウンロードしてインストールする"""
    import subprocess
    import urllib.request

    # Tesseractのtessdataディレクトリを特定
    try:
        output = subprocess.check_output(
            ["tesseract", "--print-parameters"],
            stderr=subprocess.STDOUT,
            text=True,
        )
    except Exception:
        raise HTTPException(500, "Tesseract OCRが見つかりません")

    # tessdata ディレクトリを探す
    tessdata_dir = None
    # 一般的なパスを試行
    candidates = []
    try:
        # tesseract --print-parameters の出力からtessdata_prefixを取得
        for line in output.split("\n"):
            if "tessdata" in line.lower():
                parts = line.strip().split()
                for p in parts:
                    if os.path.isdir(p):
                        candidates.append(p)
    except Exception:
        pass

    # 一般的なインストール先を追加
    if sys.platform == "win32":
        candidates += [
            r"C:\Program Files\Tesseract-OCR\tessdata",
            r"C:\Program Files (x86)\Tesseract-OCR\tessdata",
        ]
    else:
        candidates += [
            "/usr/share/tesseract-ocr/5/tessdata",
            "/usr/share/tesseract-ocr/4.00/tessdata",
            "/usr/share/tessdata",
            "/usr/local/share/tessdata",
        ]

    for c in candidates:
        if os.path.isdir(c):
            tessdata_dir = c
            break

    if not tessdata_dir:
        raise HTTPException(500, "tessdataディレクトリが見つかりません")

    jpn_path = os.path.join(tessdata_dir, "jpn.traineddata")
    if os.path.exists(jpn_path):
        return {"message": "日本語データは既にインストール済みです", "path": jpn_path}

    # GitHubからダウンロード
    url = "https://github.com/tesseract-ocr/tessdata_best/raw/main/jpn.traineddata"
    try:
        urllib.request.urlretrieve(url, jpn_path)
    except Exception as e:
        raise HTTPException(500, f"ダウンロードに失敗しました: {e}")

    return {"message": "日本語OCRデータをインストールしました", "path": jpn_path}


def open_browser(port: int):
    """ブラウザを自動で開く"""
    webbrowser.open(f"http://localhost:{port}")


def main():
    """エントリーポイント"""
    import uvicorn

    port = config["port"]

    if config["auto_open_browser"]:
        Timer(1.5, open_browser, args=[port]).start()

    print(f"PDF変換ツール を起動しています...")
    print(f"ブラウザで http://localhost:{port} を開いてください")
    print("終了するには Ctrl+C を押してください")

    uvicorn.run(app, host="127.0.0.1", port=port, log_level="warning")


if __name__ == "__main__":
    main()
